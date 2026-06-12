"""Generate the Word explanation document for the final project."""

import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def add_paragraph(doc, text, bold=False, italic=False, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    return p


def main():
    doc = Document()

    # Title
    title = doc.add_heading("AI 课程期末大作业说明文档", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Deep Research Agent — 基于 LangGraph 的智能研究助手")
    run.font.size = Pt(14)
    run.bold = True

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = info.add_run("作者：Cran\n杭州电子科技大学 英语专业")
    run.font.size = Pt(10.5)

    doc.add_paragraph()

    # 1. 项目概述
    add_heading(doc, "一、项目概述", level=1)
    add_paragraph(doc,
        "本项目是人工智能课程期末大作业，要求使用 LangGraph 编制一个功能自定、逻辑闭环的 AI Agent。 "
        "我选择实现一个『深度研究助手』（Deep Research Agent）：用户输入一个研究主题，Agent 自动完成 "
        "主题分析、资料搜索、大纲生成、报告撰写、质量反思与终稿输出，最终生成 Markdown 与 Word 研究报告。")

    # 2. 功能说明
    add_heading(doc, "二、功能说明", level=1)
    functions = [
        ("1. 主题分析（analyze_topic）", "提炼研究问题、核心概念与推荐搜索词。"),
        ("2. 资料搜索（research）", "调用 TokenDance UniFuncs web-search 实时搜索网络资料。"),
        ("3. 大纲生成（generate_outline）", "根据主题与搜索结果生成结构化报告大纲。"),
        ("4. 报告撰写（draft_report）", "基于大纲与资料撰写报告正文。"),
        ("5. 质量反思（reflect）", "评估报告质量，指出优缺点并给出改进建议。"),
        ("6. 终稿输出（finalize）", "润色整合，输出最终 Markdown 与 Word 文档。"),
    ]
    for name, desc in functions:
        p = add_paragraph(doc, "", size=10.5)
        r = p.add_run(name)
        r.bold = True
        p.add_run("：" + desc)

    # 3. 技术栈
    add_heading(doc, "三、技术栈", level=1)
    add_paragraph(doc, "本作业采用 2025–2026 年较新的技术链：")
    stack = [
        "LangGraph 1.2.4：Agent 状态图编排",
        "langchain_openai 1.3.0：通过 OpenAI 兼容接口调用大模型",
        "TokenDance（词元跳动）：云端大模型网关，主模型 qwen3.7-plus",
        "UniFuncs web-search / web-reader：实时搜索与网页阅读工具",
        "Gradio 6.18.0：Web 用户界面",
        "python-docx 1.2.0：Word 文档导出",
    ]
    for item in stack:
        add_paragraph(doc, item, size=10.5).style = "List Bullet"

    # 4. 作业要求对照
    add_heading(doc, "四、作业要求对照", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "作业要求"
    hdr_cells[1].text = "实现情况"
    hdr_cells[2].text = "说明"

    rows = [
        ("使用 LangGraph", "✅ 已实现", "agent/graph.py 完整 StateGraph"),
        ("逻辑闭环", "✅ 已实现", "主题→搜索→大纲→起草→反思→终稿"),
        ("云端大模型通信", "✅ 已实现", "TokenDance qwen3.7-plus"),
        ("≥3 种 MessageState", "✅ 已实现", "Human/AIMessage/System/ToolMessage"),
        ("≥4 个功能节点", "✅ 已实现", "共 6 个节点"),
        ("≥1 条 Loop/Concurrency 边", "✅ 已实现", "reflect→draft_report 条件循环"),
        ("复杂度≥Drafter Agent", "✅ 已实现", "节点更多、流程更长"),
        (".ipynb 展示执行结果", "✅ 已实现", "notebook/final_project.ipynb"),
        ("Word 说明文档", "✅ 已实现", "docs/说明文档.docx"),
    ]
    for req, status, note in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = req
        row_cells[1].text = status
        row_cells[2].text = note

    doc.add_paragraph()

    # 5. 项目结构
    add_heading(doc, "五、项目结构", level=1)
    structure = """
final/
├── agent/              # LangGraph Agent 核心代码
│   ├── state.py        # 状态定义
│   ├── llm.py          # TokenDance LLM 配置
│   ├── tools.py        # UniFuncs 搜索/阅读工具
│   ├── nodes.py        # 6 个功能节点
│   ├── graph.py        # LangGraph 图编译
│   └── export.py       # Word 导出
├── app/                # Gradio Web UI
│   └── main.py
├── notebook/           # 作业提交 Notebook
│   └── final_project.ipynb
├── docs/               # 说明文档与调研笔记
│   ├── 说明文档.docx
│   └── research_notes.md
├── output/             # 生成的报告文件
├── requirements.txt
├── pyproject.toml
└── README.md
"""
    add_paragraph(doc, structure.strip(), size=9)

    # 6. 运行方式
    add_heading(doc, "六、运行方式", level=1)
    steps = [
        "配置环境变量：复制 .env.example 为 .env，填入 TOKENDANCE_API_KEY。",
        "安装依赖：pip install -r requirements.txt",
        "运行 Notebook：jupyter notebook notebook/final_project.ipynb",
        "启动 Web UI：python -m app.main，浏览器访问 http://服务器IP:7860",
    ]
    for i, step in enumerate(steps, 1):
        add_paragraph(doc, f"{i}. {step}", size=10.5).style = "List Number"

    # 7. 创新与特色
    add_heading(doc, "七、创新与特色", level=1)
    features = [
        "使用 2025–2026 年最新的 LangGraph 1.x 和 Gradio 6.x 构建。",
        "接入 TokenDance 多模型网关，使用 qwen3.7-plus 进行研究与生成。",
        "使用 UniFuncs 实时搜索与网页阅读工具，替代传统 DuckDuckGo。",
        "实现反思-迭代循环，自动提升报告质量。",
        "提供精美的 Gradio Web UI，支持流式过程展示与 Word 下载。",
    ]
    for f in features:
        add_paragraph(doc, f, size=10.5).style = "List Bullet"

    # 8. 总结
    add_heading(doc, "八、总结", level=1)
    add_paragraph(doc,
        "本项目完成了一个功能完整、复杂度高于参考 Drafter Agent 的 LangGraph AI Agent。 "
        "Agent 能够与云端大模型通信，使用多种 MessageState 类型，包含 6 个功能节点和 1 条反思循环边， "
        "并提供了可部署的 Gradio Web UI。所有代码模块化、可维护，符合期末大作业的提交要求。")

    output_path = os.path.join(os.path.dirname(__file__), "说明文档.docx")
    doc.save(output_path)
    print(f"说明文档已保存到: {os.path.abspath(output_path)}")


if __name__ == "__main__":
    main()
